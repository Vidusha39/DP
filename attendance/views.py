from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import HttpResponse
from django.utils import timezone
from datetime import date, timedelta
from .models import SundaySession, AttendanceRecord
from academics.models import ClassSection, AcademicYear
from students.models import Student
from teachers.models import TeacherAllocation
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def get_most_recent_sunday():
    today = date.today()
    # In Python, Monday is 0, Sunday is 6
    idx = (today.weekday() + 1) % 7
    sun = today - timedelta(days=idx)
    return sun


@login_required
def mark_attendance(request, class_id):
    """
    Mobile-optimized Sunday attendance marking interface for teachers.
    """
    class_section = get_object_or_404(
        ClassSection.objects.select_related('grade_level', 'academic_year', 'classroom'), 
        id=class_id
    )
    user = request.user
    is_staff_or_admin = user.is_superuser or (hasattr(user, 'profile') and user.profile.role in ['ADMIN', 'PRINCIPAL'])

    # Permission check
    if not is_staff_or_admin:
        teacher_profile = getattr(user, 'teacher_profile', None)
        is_allocated = TeacherAllocation.objects.filter(
            teacher=teacher_profile,
            class_section=class_section
        ).exists()
        if not is_allocated:
            messages.error(request, 'මෙම පන්තියේ පැමිණීම සටහන් කිරීමට ඔබට අවසර නොමැත.')
            return redirect('dashboard')

    # Selected or default Sunday
    date_str = request.GET.get('date')
    if date_str:
        try:
            selected_date = date.fromisoformat(date_str)
        except ValueError:
            selected_date = get_most_recent_sunday()
    else:
        selected_date = get_most_recent_sunday()

    # Ensure SundaySession exists
    session, _ = SundaySession.objects.get_or_create(
        date=selected_date,
        defaults={'academic_year': class_section.academic_year, 'term': 1}
    )

    # Active students in this class
    students = Student.objects.filter(
        assigned_class=class_section,
        status='APPROVED'
    ).order_by('admission_number', 'name_with_initials_sinhala')

    # Handle POST (saving attendance)
    if request.method == 'POST':
        for student in students:
            status_val = request.POST.get(f'status_{student.id}', 'PRESENT')
            AttendanceRecord.objects.update_or_create(
                session=session,
                student=student,
                defaults={
                    'class_section': class_section,
                    'status': status_val,
                    'marked_by': user,
                }
            )
        messages.success(request, f'{class_section.grade_level.name} - {selected_date.strftime("%Y-%m-%d")} දින පැමිණීම සාර්ථකව සුරැකිණි.')
        return redirect(f'/attendance/mark/{class_id}/?date={selected_date.isoformat()}')

    # Existing records for this session
    existing_records = {
        rec.student_id: rec.status 
        for rec in AttendanceRecord.objects.filter(session=session, class_section=class_section)
    }

    students_with_status = []
    for s in students:
        status = existing_records.get(s.id, 'PRESENT')
        students_with_status.append({'student': s, 'status': status})

    return render(request, 'attendance/mark_attendance.html', {
        'class_section': class_section,
        'selected_date': selected_date,
        'session': session,
        'students_with_status': students_with_status,
        'total_students': students.count(),
    })


@login_required
def attendance_overview(request):
    """
    Overview of attendance reports.
    """
    user = request.user
    is_staff_or_admin = user.is_superuser or (hasattr(user, 'profile') and user.profile.role in ['ADMIN', 'PRINCIPAL'])

    current_year = AcademicYear.objects.filter(is_current=True).first()

    if is_staff_or_admin:
        classes = ClassSection.objects.filter(academic_year=current_year).select_related('grade_level')
    else:
        teacher_profile = getattr(user, 'teacher_profile', None)
        allocated_ids = TeacherAllocation.objects.filter(
            teacher=teacher_profile
        ).values_list('class_section_id', flat=True)
        classes = ClassSection.objects.filter(id__in=allocated_ids).select_related('grade_level')

    recent_sessions = SundaySession.objects.all().order_by('-date')[:10]

    return render(request, 'attendance/overview.html', {
        'classes': classes,
        'recent_sessions': recent_sessions,
        'current_year': current_year,
    })


@login_required
def export_excel(request, class_id):
    """
    Generates an Excel spreadsheet (.xlsx) with student roster and parent contacts.
    Full Sinhala Unicode text preserved.
    """
    class_section = get_object_or_404(
        ClassSection.objects.select_related('grade_level', 'academic_year', 'classroom'), 
        id=class_id
    )

    students = Student.objects.filter(
        assigned_class=class_section,
        status='APPROVED'
    ).select_related('guardian').order_by('admission_number', 'name_with_initials_sinhala')

    # Create Workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{class_section.grade_level.name}"

    # Header Styling
    title_font = Font(name='Arial', size=16, bold=True, color='800000') # Maroon
    header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
    header_fill = Font(name='Arial', size=11, bold=True)
    maroon_fill = PatternFill(start_color='800000', end_color='800000', fill_type='solid')
    thin_border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )

    # Title rows
    ws.merge_cells('A1:H1')
    ws['A1'] = "දහම් පාසල - ශිෂ්‍ය නාම ලේඛනය (Class Roster)"
    ws['A1'].font = title_font
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')

    ws.merge_cells('A2:H2')
    ws['A2'] = f"ශ්‍රේණිය: {class_section.grade_level.name} | අධ්‍යයන වර්ෂය: {class_section.academic_year.year} | මුළු සිසුන්: {students.count()}"
    ws['A2'].font = Font(name='Arial', size=11, italic=True)
    ws['A2'].alignment = Alignment(horizontal='center')

    # Column Headers (Sinhala)
    headers = [
        "අනු අංකය", 
        "ඇතුළත් වීමේ අංකය", 
        "ශිෂ්‍යයාගේ නම", 
        "ස්ත්‍රී / පුරුෂ", 
        "උපන් දිනය", 
        "දෙමාපිය / භාරකරුගේ නම", 
        "දුරකථන අංකය", 
        "ලිපිනය"
    ]

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = maroon_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # Rows
    for row_idx, student in enumerate(students, start=5):
        ws.cell(row=row_idx, column=1, value=row_idx - 4).alignment = Alignment(horizontal='center')
        ws.cell(row=row_idx, column=2, value=student.admission_number or "-").alignment = Alignment(horizontal='center')
        ws.cell(row=row_idx, column=3, value=student.name_with_initials_sinhala)
        ws.cell(row=row_idx, column=4, value=student.get_gender_display()).alignment = Alignment(horizontal='center')
        ws.cell(row=row_idx, column=5, value=student.dob.strftime('%Y-%m-%d') if student.dob else "-").alignment = Alignment(horizontal='center')
        ws.cell(row=row_idx, column=6, value=student.guardian.full_name_sinhala)
        ws.cell(row=row_idx, column=7, value=student.guardian.primary_phone).alignment = Alignment(horizontal='center')
        ws.cell(row=row_idx, column=8, value=student.guardian.address)

        for col_idx in range(1, 9):
            ws.cell(row=row_idx, column=col_idx).border = thin_border

    # Adjust Column Widths
    col_widths = [10, 18, 32, 14, 14, 30, 16, 40]
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f"Roster_{class_section.grade_level.english_name or 'Class'}_{class_section.academic_year.year}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


@login_required
def print_roster(request, class_id):
    """
    Clean printable A4 class roster view for manual marking and paper filing.
    """
    class_section = get_object_or_404(
        ClassSection.objects.select_related('grade_level', 'academic_year', 'classroom'), 
        id=class_id
    )
    students = Student.objects.filter(
        assigned_class=class_section,
        status='APPROVED'
    ).select_related('guardian').order_by('admission_number', 'name_with_initials_sinhala')

    allocations = class_section.teacher_allocations.select_related('teacher')

    return render(request, 'attendance/print_roster.html', {
        'class_section': class_section,
        'students': students,
        'allocations': allocations,
        'print_date': date.today(),
    })
